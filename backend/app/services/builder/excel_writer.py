"""Stage 4 — Excel writer.

Strategi: kerja di atas SALINAN file yang di-upload (PAKEM terjaga). Untuk tiap
item, tulis ke baris aslinya (`excel_row`):
  - col G (harga satuan) = referensi ke sheet 'Resume Analisa' (HSP final terkalibrasi)
  - col H (jumlah)       = G * E (volume)   ← bukan range, jadi tak ada double-count
  - col I (TKDN)         = referensi ke 'Resume Analisa' col E (faktor TKDN)
Lalu buat/refresh sheet 'Resume Analisa' (cols A-I sesuai spesifikasi build.md).

Decoupled dari DB: input berupa `PricedItemRecord` (plain dataclass).
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import openpyxl.reader.drawings as drw  # type: ignore
from loguru import logger
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from app.services.builder.formulas import cell_ref, eq
from app.services.parser.excel_parser import ColumnMap, normalize_text

# Patch openpyxl agar tahan file dengan broken images (sama seperti parser).
drw.find_images = lambda a, t: ([], [])

RESUME_SHEET = "Resume Analisa"
# Header Resume Analisa (Known Issue: A=NO B=URAIAN C=SAT D=HARGA E=TKDN F=TIPE G=KODE H=SUMBER I=TIER).
RESUME_HEADERS = ["NO", "URAIAN", "SAT", "HARGA", "NILAI TKDN", "TIPE", "KODE", "SUMBER", "TIER"]
RESUME_HEADER_ROW = 1
RESUME_DATA_START = 2


@dataclass
class PricedItemRecord:
    """Satu item paket beserta hasil match + harga (untuk ditulis ke Excel)."""

    sheet_name: str
    excel_row: int
    norm_uraian: str
    norm_satuan: str
    uraian: str
    satuan: str
    volume: float
    harga: float  # final_hsp (terkalibrasi)
    tkdn: float
    tipe: str  # ahsp | lumpsum | unresolved
    kode: str
    sumber: str
    tier: str

    @property
    def key(self) -> tuple[str, str]:
        return (self.norm_uraian, self.norm_satuan)


def _build_resume_sheet(wb, records: list[PricedItemRecord]) -> dict[tuple[str, str], int]:
    """Buat/refresh sheet Resume Analisa. Return map unique_key -> nomor baris."""
    if RESUME_SHEET in wb.sheetnames:
        del wb[RESUME_SHEET]
    ws = wb.create_sheet(RESUME_SHEET)

    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="E8EEF7")
    for col_idx, head in enumerate(RESUME_HEADERS, start=1):
        c = ws.cell(row=RESUME_HEADER_ROW, column=col_idx, value=head)
        c.font = bold
        c.fill = fill

    # Satu baris per unique key, urut deterministik (sheet lalu uraian).
    seen: dict[tuple[str, str], int] = {}
    uniques: list[PricedItemRecord] = []
    for r in sorted(records, key=lambda x: (x.sheet_name, x.uraian)):
        if r.key not in seen:
            seen[r.key] = -1
            uniques.append(r)

    row_map: dict[tuple[str, str], int] = {}
    for i, r in enumerate(uniques):
        row = RESUME_DATA_START + i
        ws.cell(row=row, column=1, value=i + 1)  # NO
        ws.cell(row=row, column=2, value=r.uraian)  # URAIAN
        ws.cell(row=row, column=3, value=r.satuan)  # SAT
        ws.cell(row=row, column=4, value=round(r.harga, 2))  # HARGA
        ws.cell(row=row, column=5, value=round(r.tkdn, 4))  # NILAI TKDN (faktor 0-1)
        ws.cell(row=row, column=6, value=r.tipe.upper())  # TIPE
        ws.cell(row=row, column=7, value=r.kode)  # KODE
        ws.cell(row=row, column=8, value=r.sumber)  # SUMBER
        ws.cell(row=row, column=9, value=r.tier)  # TIER
        row_map[r.key] = row

    # Lebar kolom enak dibaca.
    widths = {"A": 5, "B": 50, "C": 8, "D": 14, "E": 10, "F": 10, "G": 16, "H": 36, "I": 6}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    return row_map


def _inject_paket_pricing(
    wb,
    records: list[PricedItemRecord],
    row_map: dict[tuple[str, str], int],
    col_map: ColumnMap | None = None,
) -> int:
    """Tulis formula harga/jumlah/TKDN ke baris asli tiap item. Return jumlah ditulis."""
    cm = col_map or ColumnMap()
    written = 0
    for r in records:
        if r.sheet_name not in wb.sheetnames:
            continue
        ws = wb[r.sheet_name]
        resume_row = row_map.get(r.key)
        if resume_row is None:
            continue
        gcell = f"{cm.harga}{r.excel_row}"
        ecell = f"{cm.vol}{r.excel_row}"
        # G = harga satuan -> ref Resume Analisa D
        ws[gcell] = eq(cell_ref(RESUME_SHEET, f"D{resume_row}"))
        # H = jumlah = G * E (volume). Bukan SUM range -> aman double-count.
        ws[f"{cm.jumlah}{r.excel_row}"] = eq(f"{gcell}*{ecell}")
        # I = TKDN -> ref Resume Analisa E
        ws[f"{cm.tkdn}{r.excel_row}"] = eq(cell_ref(RESUME_SHEET, f"E{resume_row}"))
        written += 1
    return written


def generate_workbook(
    input_path: Path | str,
    output_path: Path | str,
    records: list[PricedItemRecord],
    col_map: ColumnMap | None = None,
) -> dict:
    """Hasilkan workbook BOQ dari salinan file input + data harga.

    Return ringkasan {resume_rows, items_written, output_path}.
    """
    input_path = Path(input_path)
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = load_workbook(input_path)  # data_only=False: pertahankan formula existing
    row_map = _build_resume_sheet(wb, records)
    written = _inject_paket_pricing(wb, records, row_map, col_map)

    wb.save(output_path)
    logger.info(
        f"Workbook generated: {output_path} "
        f"(resume_rows={len(row_map)}, items_written={written})"
    )
    return {
        "resume_rows": len(row_map),
        "items_written": written,
        "output_path": str(output_path),
    }


def normalize_key(uraian: str, satuan: str) -> tuple[str, str]:
    """Util untuk membuat unique key konsisten dengan parser."""
    return (normalize_text(uraian), normalize_text(satuan))
