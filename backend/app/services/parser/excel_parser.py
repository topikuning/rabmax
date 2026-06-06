"""Stage 1: Parser — Excel file lelang → structured data.

Deterministic parser first. AI fallback (Stage 1B) untuk struktur sheet yang
tidak terdeteksi atau confidence rendah.
"""

import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl.reader.drawings as drw  # type: ignore
from loguru import logger
from openpyxl import load_workbook

# Patch openpyxl untuk handle file dengan broken images
drw.find_images = lambda a, t: ([], [])


@dataclass
class ColumnMap:
    """Mapping kolom paket sheet."""

    no: str = "A"  # Kolom nomor (1, 2, a, b, etc)
    uraian: str = "B"
    vol: str = "E"
    sat: str = "F"
    harga: str = "G"
    jumlah: str = "H"
    tkdn: str = "I"
    kdn: str = "J"


@dataclass
class PaketSheetStructure:
    """Hasil parse 1 paket sheet."""

    sheet_name: str
    header_row: int
    max_row: int
    column_map: ColumnMap
    item_count: int


@dataclass
class ParsedItem:
    """Item individual dari paket sheet."""

    sheet_name: str
    excel_row: int
    no_label: str | None
    uraian: str
    satuan: str
    volume: float
    parent_uraian: str | None = None
    norm_uraian: str = ""
    norm_satuan: str = ""
    # Mode B (RAB terisi): harga satuan (col G) & jumlah (col H) dari file HPS.
    harga: float | None = None
    jumlah: float | None = None


@dataclass
class ParseResult:
    paket_sheets: list[PaketSheetStructure] = field(default_factory=list)
    aggregator_sheets: list[str] = field(default_factory=list)
    items: list[ParsedItem] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    needs_ai_assist: list[str] = field(
        default_factory=list,
        metadata={"description": "Sheet names yang struktur-nya tidak terdeteksi"},
    )


# Aggregator sheets — sheets dengan struktur khusus, bukan paket biasa
AGGREGATOR_SHEET_NAMES = {"REKAP", "RAB", "Sub Resume EE", "Bahan & Upah", "ANALISA", "Resume Analisa"}


def normalize_text(s: str) -> str:
    """Normalize untuk matching: lowercase, strip whitespace, unicode replace."""
    s = str(s).lower().strip()
    s = re.sub(r"\s+", " ", s)
    return (
        s.replace("½", "1/2")
        .replace("¼", "1/4")
        .replace("¾", "3/4")
        .replace("ø", "o")
        .replace("∅", "o")
    )


def detect_header_row(ws, max_check: int = 15) -> int | None:
    """Cari header row yang mengandung 'NO', 'URAIAN'/'JENIS', dan 'VOL'/'JUMLAH'."""
    for r in range(1, min(max_check, ws.max_row + 1)):
        cells = [ws.cell(row=r, column=c).value for c in range(1, 12)]
        text = " ".join(str(c or "").upper() for c in cells)
        if (
            "NO" in text
            and ("URAIAN" in text or "JENIS" in text)
            and ("VOL" in text or "JUMLAH" in text)
        ):
            return r
    return None


def is_section_header(uraian: str) -> bool:
    """Cek apakah uraian adalah section header (bukan item pekerjaan)."""
    ul = uraian.lower()
    section_keywords = [
        "area halaman",
        "elevasi ",
        "shelter pendaratan ",
        "bangunan toilet",
        "pekerjaan struktural",
        "pekerjaan arsitektural",
    ]
    return any(kw in ul for kw in section_keywords)


def find_parent_uraian(ws, current_row: int, header_row: int) -> str | None:
    """Look back ke parent row untuk dapat uraian pekerjaan jika current row hanya nama lokasi."""
    for back_r in range(current_row - 1, max(header_row, current_row - 15), -1):
        pb = ws.cell(row=back_r, column=2).value
        pe = ws.cell(row=back_r, column=5).value  # vol
        if pb and isinstance(pb, str) and len(pb) > 5 and not isinstance(pe, (int, float)):
            # Skip section headers yang ALL CAPS dan pendek (mis. "PEKERJAAN PERSIAPAN")
            if pb.isupper() and len(pb.split()) < 5:
                continue
            return str(pb)
    return None


def parse_workbook(file_path: Path | str) -> ParseResult:
    """Parse Excel file lelang dan return semua items terstruktur.

    Algoritma:
    1. Iterate semua sheet.
    2. Identifikasi paket sheet (punya header row standar) vs aggregator.
    3. Extract item: row dengan vol+sat valid.
    4. Untuk item yang uraian-nya = section header (e.g. "Area Halaman..."),
       lookback ke parent row dengan uraian asli.
    """
    result = ParseResult()
    file_path = Path(file_path)
    logger.info(f"Parsing workbook: {file_path}")

    wb = load_workbook(file_path, data_only=True)

    for sn in wb.sheetnames:
        ws = wb[sn]

        if sn in AGGREGATOR_SHEET_NAMES:
            result.aggregator_sheets.append(sn)
            continue

        header_row = detect_header_row(ws)
        if header_row is None:
            result.warnings.append(
                f"Sheet '{sn}': header row not detected, marking for AI assist"
            )
            result.needs_ai_assist.append(sn)
            continue

        # Default column map (standar Indonesian BOQ template)
        col_map = ColumnMap()

        item_count = 0
        for r in range(header_row + 1, ws.max_row + 1):
            b = ws.cell(row=r, column=2).value  # uraian
            e = ws.cell(row=r, column=5).value  # vol
            f = ws.cell(row=r, column=6).value  # sat
            a = ws.cell(row=r, column=1).value  # no

            if (
                isinstance(e, (int, float))
                and e > 0
                and b
                and isinstance(b, str)
                and len(b) > 3
                and f
            ):
                uraian = str(b).strip()
                satuan = str(f).strip()

                parent = None
                if is_section_header(uraian):
                    parent = find_parent_uraian(ws, r, header_row)
                    effective_uraian = parent if parent else uraian
                else:
                    effective_uraian = uraian

                result.items.append(
                    ParsedItem(
                        sheet_name=sn,
                        excel_row=r,
                        no_label=str(a) if a is not None else None,
                        uraian=uraian,
                        satuan=satuan,
                        volume=float(e),
                        parent_uraian=parent,
                        norm_uraian=normalize_text(effective_uraian),
                        norm_satuan=normalize_text(satuan),
                    )
                )
                item_count += 1

        result.paket_sheets.append(
            PaketSheetStructure(
                sheet_name=sn,
                header_row=header_row,
                max_row=ws.max_row,
                column_map=col_map,
                item_count=item_count,
            )
        )

    logger.info(
        f"Parsed {len(result.items)} items from {len(result.paket_sheets)} paket sheets. "
        f"Aggregator sheets: {len(result.aggregator_sheets)}. "
        f"AI assist needed: {len(result.needs_ai_assist)}"
    )
    return result


def parse_filled_rab(file_path: Path | str) -> ParseResult:
    """Parse RAB terisi (Mode B — profit analysis).

    Sama seperti parse_workbook tapi juga menangkap col G (harga satuan) dan
    col H (jumlah) sebagai nilai HPS dari file terisi.
    """
    result = parse_workbook(file_path)

    wb = load_workbook(Path(file_path), data_only=True)
    # Index item per (sheet, row) untuk overlay harga/jumlah.
    by_loc = {(it.sheet_name, it.excel_row): it for it in result.items}

    for sn in wb.sheetnames:
        if sn in AGGREGATOR_SHEET_NAMES:
            continue
        ws = wb[sn]
        for (s_name, r), item in list(by_loc.items()):
            if s_name != sn:
                continue
            g = ws.cell(row=r, column=7).value  # harga satuan
            h = ws.cell(row=r, column=8).value  # jumlah
            if isinstance(g, (int, float)):
                item.harga = float(g)
            if isinstance(h, (int, float)):
                item.jumlah = float(h)
            # Fallback hitung jumlah kalau hanya harga yang ada.
            if item.jumlah is None and item.harga is not None:
                item.jumlah = round(item.harga * item.volume, 2)

    return result
