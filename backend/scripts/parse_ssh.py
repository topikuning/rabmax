"""Parser SSH per-kota — dari 'Daftar_Upah_Bahan' (harga satuan resmi per kota/kab).

Tiap file Excel berisi section "DAFTAR HARGA SATUAN {UPAH PEKERJA|BAHAN BANGUNAN|
ALAT BERAT/SEWA ...}" dengan header 3-baris (judul / daftar kota / provinsi), lalu
tabel NO. | JENIS | SATUAN | HARGA. Satu section bisa untuk BEBERAPA kota (dipisah
koma) → di-expand jadi satu baris harga per kota.

Output: bahan_upah_ssh_2026.jsonl.gz (format seed_bahan_upah; tiap item punya
provinsi + kota + category + harga). Provinsi fallback dari nama folder.

Jalankan:
    cd backend
    python -m scripts.parse_ssh /path/to/Daftar_Upah_Bahan [out_file.jsonl.gz]
"""

from __future__ import annotations

import gzip
import json
import re
import sys
from pathlib import Path

import openpyxl

TAHUN = 2026
SECTION_CAT = [
    ("UPAH", "upah"), ("BAHAN", "bahan"), ("MATERIAL", "bahan"),
    ("SEWA", "alat"), ("ALAT", "alat"), ("PERALATAN", "alat"),
]


def _clean(s: object) -> str:
    return re.sub(r"\s+", " ", str(s or "").replace("\n", " ")).strip()


def _num(v: object) -> float | None:
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        s = v.strip().replace("Rp", "").replace(" ", "")
        if "," in s and "." in s:
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
        try:
            return float(s)
        except ValueError:
            return None
    return None


def _category(title: str) -> str:
    up = title.upper()
    for key, cat in SECTION_CAT:
        if key in up:
            return cat
    return "bahan"


# Satuan upah beragam (Orang/Hari, org/hari, Hari, OH, …) → kanonik 'OH'
# agar cocok dengan komponen AHSP (satuan 'OH').
_OH_RE = re.compile(r"(orang\s*/?\s*hari|org\s*/?\s*hari|^hari$|o\.?h)", re.I)


def _norm_satuan(satuan: str, category: str) -> str:
    if category == "upah" and (_OH_RE.search(satuan) or not satuan):
        return "OH"
    return satuan


def _norm_prov(prov: str) -> str:
    """Provinsi → Title Case rapi (samakan 'LAMPUNG' & folder 'Lampung')."""
    p = re.sub(r"\s+", " ", str(prov or "").strip()).title()
    return p.replace("Dki ", "DKI ").replace("Diy", "DIY")


def _parse_header(text: str, folder_prov: str) -> tuple[str, list[str], str]:
    """Header multi-baris → (category, [kota...], provinsi)."""
    parts = [p.strip() for p in str(text).split("\n") if p.strip()]
    cat = _category(parts[0] if parts else "")
    kota: list[str] = []
    prov = folder_prov
    for p in parts[1:]:
        up = p.upper()
        if up.startswith("PROVINSI") or up.startswith("PROV."):
            prov = _clean(re.sub(r"^PROV(INSI)?\.?", "", p, flags=re.I))
        elif any(k in up for k in ("KOTA", "KAB", "KABUPATEN")):
            # daftar kota dipisah koma
            for tok in p.split(","):
                t = _clean(tok)
                if t:
                    kota.append(t)
    return cat, kota, prov


def _is_table_header(row: list) -> int | None:
    """Return index kolom HARGA bila baris ini header tabel (NO|JENIS|SATUAN|HARGA)."""
    has_no = has_harga = False
    harga_i = None
    for i, c in enumerate(row):
        s = _clean(c).upper().rstrip(".")
        if s in ("NO", "NO."):
            has_no = True
        if s.startswith("HARGA"):
            has_harga = True
            harga_i = i
    return harga_i if (has_no and has_harga) else None


def parse_file(path: Path, folder_prov: str) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out: list[dict] = []
    for sn in wb.sheetnames:
        rows = [list(r) for r in wb[sn].iter_rows(values_only=True)]
        cat, kota, prov = "bahan", [], folder_prov
        harga_i = nama_i = sat_i = None
        for row in rows:
            # Section header (sel berisi 'DAFTAR HARGA SATUAN ...').
            head_cell = next((c for c in row if isinstance(c, str) and "DAFTAR HARGA SATUAN" in c.upper()), None)
            if head_cell:
                cat, kota, prov = _parse_header(head_cell, folder_prov)
                harga_i = None  # tunggu header tabel baru
                continue
            ti = _is_table_header(row)
            if ti is not None:
                harga_i = ti
                # nama = kolom teks ke-2 (setelah NO); satuan = sebelum harga.
                nama_i = 1
                for i, c in enumerate(row):
                    s = _clean(c).upper()
                    if s.startswith("JENIS") or s.startswith("URAIAN"):
                        nama_i = i
                    if s == "SATUAN":
                        sat_i = i
                continue
            if harga_i is None:
                continue
            nama = _clean(row[nama_i]) if nama_i is not None and nama_i < len(row) else ""
            harga = _num(row[harga_i]) if harga_i < len(row) else None
            if not nama or harga is None or harga <= 0 or nama.upper() in ("NO", "JENIS"):
                continue
            satuan = _norm_satuan(_clean(row[sat_i]) if sat_i is not None and sat_i < len(row) else "", cat)
            prov_c = _norm_prov(prov)
            targets = kota or [None]  # bila tak ada kota → level provinsi
            for kt in targets:
                out.append({"nama": nama, "satuan": satuan, "harga": harga,
                            "category": cat, "provinsi": prov_c, "kota": kt,
                            "tier": "A", "tahun": TAHUN,
                            "source_label": f"SSH {kt or prov_c} {TAHUN}"})
    wb.close()
    return out


_PROV_RE = re.compile(r"^\d+\s+PROVINSI\s+", re.I)


def build(root: str) -> list[dict]:
    base = Path(root)
    items: list[dict] = []
    for xlsx in sorted(base.rglob("*.xlsx")):
        if xlsx.name.startswith("~$"):
            continue
        folder_prov = _PROV_RE.sub("", xlsx.parent.name).strip().title()
        try:
            items.extend(parse_file(xlsx, folder_prov))
        except Exception as e:  # noqa: BLE001
            print(f"  ! gagal {xlsx.name}: {e}")
    return items


def main() -> None:
    if len(sys.argv) < 2:
        print("Usage: python -m scripts.parse_ssh <Daftar_Upah_Bahan dir> [out.jsonl.gz]")
        raise SystemExit(2)
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else \
        Path(__file__).resolve().parent.parent / "seed_data" / "bahan_upah_ssh_2026.jsonl.gz"
    items = build(sys.argv[1])
    from collections import Counter
    by_cat = Counter(i["category"] for i in items)
    by_prov = Counter(i["provinsi"] for i in items)
    print(f"Parsed {len(items)} baris harga; kategori={dict(by_cat)}; provinsi={len(by_prov)}")
    lines = [json.dumps({"meta": {"source_label": f"SSH {TAHUN}", "tahun": TAHUN, "tier": "A"}}, ensure_ascii=False)]
    lines += [json.dumps(it, ensure_ascii=False) for it in items]
    out.write_bytes(gzip.compress(("\n".join(lines) + "\n").encode("utf-8")))
    print(f"  ✓ {out.name}: {len(items)} baris ({out.stat().st_size // 1024} KB)")


if __name__ == "__main__":
    main()
