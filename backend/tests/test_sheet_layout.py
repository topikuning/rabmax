"""Test deteksi struktur dinamis (RAB beda-beda tiap proyek) + double-count."""

from openpyxl import Workbook

from app.services.parser.excel_parser import (
    analyze_sheet_layout,
    classify_row,
    detect_header_row,
)
from app.services.validator import check_workbook_double_count


def _build_dynamic_rab(path, *, double_count: bool):
    """Workbook dengan 2 seksi, masing-masing subtotal, lalu grand total.

    Geometri sengaja 'acak' (mulai row 3, jumlah item beda) untuk membuktikan
    deteksi tidak bergantung nomor baris tetap.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Paket Jalan"
    ws["A3"], ws["B3"], ws["E3"], ws["F3"], ws["H3"] = "NO", "URAIAN", "VOL", "SAT", "JUMLAH"
    # Seksi 1: 2 item (row 4-5), subtotal row 6.
    ws["B4"], ws["E4"], ws["F4"], ws["H4"] = "Galian", 10, "m3", "=G4*E4"
    ws["B5"], ws["E5"], ws["F5"], ws["H5"] = "Urugan", 5, "m3", "=G5*E5"
    ws["B6"], ws["H6"] = "Sub Jumlah", "=SUM(H4:H5)"
    # Seksi 2: 1 item (row 7), subtotal row 8.
    ws["B7"], ws["E7"], ws["F7"], ws["H7"] = "Aspal", 100, "m2", "=G7*E7"
    ws["B8"], ws["H8"] = "Sub Jumlah", "=SUM(H7:H7)"
    # Grand total row 9.
    if double_count:
        # BUG: range mencakup baris subtotal 6 & 8 → double count.
        ws["B9"], ws["H9"] = "Jumlah Total", "=SUM(H4:H8)"
    else:
        # BENAR: jumlah subtotal saja.
        ws["B9"], ws["H9"] = "Jumlah Total", "=H6+H8"
    wb.save(path)
    return ws


def test_detect_header_and_layout(tmp_path):
    p = tmp_path / "rab.xlsx"
    _build_dynamic_rab(p, double_count=False)
    from openpyxl import load_workbook

    ws = load_workbook(p)["Paket Jalan"]
    hr = detect_header_row(ws)
    assert hr == 3
    layout = analyze_sheet_layout(ws, hr)
    assert layout.item_rows == [4, 5, 7]
    assert layout.subtotal_rows == [6, 8]
    assert layout.total_row == 9
    assert len(layout.sections) == 2
    assert layout.sections[0].item_rows == [4, 5]
    assert layout.sections[0].subtotal_row == 6


def test_classify_row(tmp_path):
    p = tmp_path / "rab.xlsx"
    _build_dynamic_rab(p, double_count=False)
    from openpyxl import load_workbook

    from app.services.parser.excel_parser import ColumnMap

    ws = load_workbook(p)["Paket Jalan"]
    cm = ColumnMap()
    assert classify_row(ws, 4, cm) == "item"
    assert classify_row(ws, 6, cm) == "subtotal"
    assert classify_row(ws, 9, cm) == "total"


def test_double_count_flagged(tmp_path):
    p = tmp_path / "bug.xlsx"
    _build_dynamic_rab(p, double_count=True)
    rep = check_workbook_double_count(p)
    assert any("double-count" in w for w in rep.warnings)


def test_clean_no_double_count(tmp_path):
    p = tmp_path / "ok.xlsx"
    _build_dynamic_rab(p, double_count=False)
    rep = check_workbook_double_count(p)
    assert not any("double-count" in w for w in rep.warnings)
