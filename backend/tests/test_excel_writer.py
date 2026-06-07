"""End-to-end test Stage 4 writer + Stage 6 validator (pakai workbook sintetis)."""

from openpyxl import Workbook, load_workbook

from app.services.builder.excel_writer import (
    RESUME_SHEET,
    SUMBER_SHEET,
    PricedItemRecord,
    generate_workbook,
    normalize_key,
)
from app.services.validator import check_records, check_workbook_formulas


def _make_input(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "PaketA"
    ws["A4"], ws["B4"], ws["E4"], ws["F4"] = "NO", "URAIAN", "VOL", "SAT"
    # Item di row 5 & 6, volume di col E.
    ws["A5"], ws["B5"], ws["E5"], ws["F5"] = 1, "Galian tanah biasa", 10.0, "m3"
    ws["A6"], ws["B6"], ws["E6"], ws["F6"] = 2, "Plesteran 1:4", 25.0, "m2"
    wb.save(path)


def _records():
    k1 = normalize_key("Galian tanah biasa", "m3")
    k2 = normalize_key("Plesteran 1:4", "m2")
    return [
        PricedItemRecord(
            "PaketA", 5, k1[0], k1[1], "Galian tanah biasa", "m3", 10.0,
            harga=50000.0, tkdn=1.0, tipe="ahsp", kode="A.1.1.1", sumber="permen_pupr_8_2023", tier="consistent",
            price_source="SSH resmi kota ×2 · Baseline nasional ×3",
        ),
        PricedItemRecord(
            "PaketA", 6, k2[0], k2[1], "Plesteran 1:4", "m2", 25.0,
            harga=80000.0, tkdn=0.95, tipe="ahsp", kode="A.4.4.1", sumber="permen_pupr_8_2023", tier="single_source",
            price_source="Baseline nasional (AHSP CK) ×4",
        ),
    ]


def test_generate_and_validate(tmp_path):
    inp = tmp_path / "in.xlsx"
    out = tmp_path / "out.xlsx"
    _make_input(inp)
    records = _records()

    summary = generate_workbook(inp, out, records)
    assert summary["resume_rows"] == 2
    assert summary["items_written"] == 2
    assert summary["sumber_rows"] == 2

    wb = load_workbook(out)
    assert RESUME_SHEET in wb.sheetnames

    # Sheet audit "Sumber Harga" berisi jejak asal harga per item.
    assert SUMBER_SHEET in wb.sheetnames
    sh = wb[SUMBER_SHEET]
    assert sh["A1"].value == "NO" and sh["G1"].value == "SUMBER HARGA"
    sumber_vals = {sh.cell(row=r, column=7).value for r in (2, 3)}
    assert "SSH resmi kota ×2 · Baseline nasional ×3" in sumber_vals
    resume = wb[RESUME_SHEET]
    # Header sesuai spesifikasi (E = NILAI TKDN).
    assert resume["A1"].value == "NO"
    assert resume["E1"].value == "NILAI TKDN"
    # Baris data: harga di D, tkdn faktor 0-1 di E.
    assert resume["D2"].value in (50000.0, 80000.0)
    assert 0.0 <= resume["E2"].value <= 1.0

    # Formula ter-inject di paket sheet.
    pa = wb["PaketA"]
    assert pa["G5"].value.startswith("=") and RESUME_SHEET in pa["G5"].value
    assert pa["H5"].value == "=G5*E5"  # jumlah = harga * volume, bukan range
    assert pa["I5"].value.startswith("=") and RESUME_SHEET in pa["I5"].value

    # Validator.
    rep = check_records(records)
    assert rep.ok
    assert rep.stats["items_priced"] == 2
    assert rep.stats["grand_total"] == 50000.0 * 10 + 80000.0 * 25

    frep = check_workbook_formulas(out, records)
    assert frep.ok
    assert frep.stats["formula_missing"] == 0


def test_validator_flags_bad_tkdn(tmp_path):
    k = normalize_key("X", "m")
    bad = [PricedItemRecord("S", 5, k[0], k[1], "X", "m", 1.0, harga=100.0, tkdn=5.0, tipe="ahsp", kode="-", sumber="-", tier="-")]
    rep = check_records(bad)
    assert not rep.ok  # TKDN 5.0 di luar 0-1
    assert any("TKDN" in e for e in rep.errors)
